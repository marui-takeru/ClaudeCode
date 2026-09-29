#!/usr/bin/env python3
"""四国の 4 空港の利用状況 (統計) と、発着する航空便 (推計) を作るスクリプト。

  python3 tools/build_airports.py

入力 (tools/src/):
  C28-21_Airport.geojson, C28-21_AirportReferencePoint.geojson
      国土数値情報「空港データ (C28、2021 年度)」(国土交通省) 商用可
  kanri_r7.xlsx   国土交通省「令和 7 年 空港管理状況調書」 (月別の乗降客数・着陸回数)
  kanri_ts.xlsx   国土交通省「暦年・年度別 空港管理状況調書 (H28〜R7)」
出力: tools/extra/airports.json   (build_network.py が自動で取り込む)

- 空港の位置・滑走路の向き・運用時間・乗降客数・着陸回数は**実データ**
- 就航路線と 1 日の便数、発着時刻は**推計** (各空港の公表情報を参考にしたおおよその値)
- 飛行機は空港から約 160km の範囲だけを描く (離陸・上昇 / 降下・着陸)
"""
import json
import math
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "src")
OUT_DIR = os.path.join(HERE, "extra")

# 空港: C28 の名称、統計の名称、着陸に使う滑走路の向き (推計: 海側から進入する向き)
AIRPORTS = [
    dict(id="MYJ", name="松山空港", c28="松山空港", stat="松山", landing="14"),
    dict(id="TAK", name="高松空港", c28="高松空港", stat="高松", landing="26"),
    dict(id="KCZ", name="高知空港", c28="高知空港", stat="高知", landing="32"),
    dict(id="TKS", name="徳島空港", c28="徳島飛行場", stat="徳島", landing="29"),
]

# 行き先の空港 (緯度, 経度)
DEST = {
    "羽田": (35.549, 139.780), "成田": (35.772, 140.393), "伊丹": (34.785, 135.438),
    "中部": (34.858, 136.805), "名古屋（小牧）": (35.255, 136.924), "福岡": (33.586, 130.451),
    "鹿児島": (31.803, 130.719), "那覇": (26.206, 127.646), "ソウル（仁川）": (37.460, 126.440),
    "釜山": (35.180, 128.940), "台北（桃園）": (25.080, 121.230), "上海（浦東）": (31.140, 121.800),
    "香港": (22.310, 113.910),
}

# 就航路線と 1 日の往復便数 (推計)。便数は目安で、曜日・季節による変動は反映しない
ROUTES = {
    "MYJ": [("羽田", 12), ("伊丹", 9), ("福岡", 4), ("成田", 2), ("中部", 2), ("鹿児島", 1), ("那覇", 1),
            ("ソウル（仁川）", 2), ("釜山", 1), ("台北（桃園）", 1), ("上海（浦東）", 1)],
    "TAK": [("羽田", 13), ("成田", 3), ("那覇", 1), ("ソウル（仁川）", 3), ("台北（桃園）", 1),
            ("上海（浦東）", 1), ("香港", 1)],
    "KCZ": [("羽田", 10), ("伊丹", 4), ("福岡", 3), ("名古屋（小牧）", 2), ("成田", 1)],
    "TKS": [("羽田", 11), ("福岡", 2)],
}
INTL = {"ソウル（仁川）", "釜山", "台北（桃園）", "上海（浦東）", "香港"}

RANGE_KM = 160        # 空港からこの距離まで描く
FINAL_KM = 15         # 最終進入の直線区間
TURN_DEG_PER_KM = 4   # 旋回の緩さ
PLANE = dict(kind="plane", cars=1, carLength=40, width=36, height=6, speed=480, dwell=0, accel=150)
COLOR = "#f5f7fa"
GROUP = dict(id="air", name="航空便（推計）")


def meters(a, b):
    lat = math.radians((a[1] + b[1]) / 2)
    return math.hypot((a[0] - b[0]) * 111320 * math.cos(lat), (a[1] - b[1]) * 110540)


def move(c, heading_deg, dist_m):
    h = math.radians(heading_deg)
    lat = c[1] + dist_m * math.cos(h) / 110540
    lon = c[0] + dist_m * math.sin(h) / (111320 * math.cos(math.radians(c[1])))
    return (lon, lat)


def bearing(a, b):
    y = math.sin(math.radians(b[0] - a[0])) * math.cos(math.radians(b[1]))
    x = math.cos(math.radians(a[1])) * math.sin(math.radians(b[1])) - \
        math.sin(math.radians(a[1])) * math.cos(math.radians(b[1])) * math.cos(math.radians(b[0] - a[0]))
    return math.degrees(math.atan2(y, x)) % 360


def turn_toward(h, target, max_deg):
    d = (target - h + 540) % 360 - 180
    return h + max(-max_deg, min(max_deg, d))


def climb_path(start, heading, dest, straight_km):
    """start から heading で straight_km 進み、行き先の方角へ旋回しながら RANGE_KM まで進む点列。"""
    pts = [start]
    c, h = start, heading
    for _ in range(int(straight_km)):
        c = move(c, h, 1000)
        pts.append(c)
    origin = start
    while meters(origin, c) < RANGE_KM * 1000:
        h = turn_toward(h, bearing(c, dest), TURN_DEG_PER_KM)
        c = move(c, h, 1000)
        pts.append(c)
    return pts


def runway_geometry(c28, refpoints, ap):
    f = next(x for x in c28["features"] if x["properties"]["C28_005"] == ap["c28"])
    p = f["properties"]
    ring = f["geometry"]["coordinates"][0]
    cx = sum(x for x, _ in ring) / len(ring)
    cy = sum(y for _, y in ring) / len(ring)
    k = math.cos(math.radians(cy))
    sxx = sum(((x - cx) * k) ** 2 for x, _ in ring)
    syy = sum((y - cy) ** 2 for _, y in ring)
    sxy = sum((x - cx) * k * (y - cy) for x, y in ring)
    axis = (90 - math.degrees(0.5 * math.atan2(2 * sxy, sxx - syy))) % 180  # 滑走路の方位 (0〜180)
    ref = refpoints[p["C28_101"].lstrip("#")]
    # 着陸方向: 滑走路番号 (例: "14" = 140°付近) に近い向き
    target = int(ap["landing"]) * 10
    heading = axis if abs((axis - target + 540) % 360 - 180) < 90 else (axis + 180) % 360
    length = int(p["C28_012"])
    threshold = move(tuple(ref), heading + 180, length / 2)
    return dict(ref=ref, heading=round(heading, 1), length=length, threshold=threshold,
                hours=[p["C28_009"], p["C28_010"]], polygon=[[round(x, 6), round(y, 6)] for x, y in ring])


def is_airport(cell, name):
    """統計表の空港名 (例: "松    山                （国管理）") が name と一致するか。"""
    return isinstance(cell, str) and re.sub(r"\s|（.*）", "", cell) == name


def stats(ap):
    """空港管理状況調書から、2025 年 (暦年) の乗降客数・着陸回数と月別、2016〜2025 年の推移を取り出す。"""
    wb = openpyxl.load_workbook(os.path.join(SRC, "kanri_ts.xlsx"), read_only=True, data_only=True)
    rows = list(wb["暦年・年度別空港管理状況調書"].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r[1] == "空港名 ：" and is_airport(r[2], ap["stat"]))
    era = {"28": 2016, "29": 2017, "30": 2018, "元": 2019}
    trend = []
    for r in rows[head + 5:head + 15]:
        y = era.get(str(r[1]), 2018 + int(r[1]) if str(r[1]).isdigit() else None)
        trend.append(dict(year=y, passengers=r[12], landings=r[4]))
    wb = openpyxl.load_workbook(os.path.join(SRC, "kanri_r7.xlsx"), read_only=True, data_only=True)
    rows = list(wb["空港管理状況調書"].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r[2] == "空港名 ：" and is_airport(r[3], ap["stat"]))
    monthly = []
    for r in rows[head + 5:head + 20]:
        if isinstance(r[2], str) and r[2].endswith("月"):
            monthly.append(dict(month=int(r[2][:-1]), passengers=r[13], landings=r[5]))
        if len(monthly) == 12:
            break
    last = trend[-1]
    total = next(r for r in rows[head:head + 40] if isinstance(r[2], str) and r[2].startswith("暦年"))
    return dict(year=2025, passengers=last["passengers"], landings=last["landings"],
                domestic=total[12], international=total[9], monthly=monthly, trend=trend)


def schedule(n, hours, offset_min):
    """運用時間内に n 便を均等に割り当てる (分単位の発時刻)。"""
    open_min = int(hours[0][:2]) * 60 + int(hours[0][2:]) + 20
    close_min = int(hours[1][:2]) * 60 + int(hours[1][2:]) - 30
    span = close_min - open_min
    out = []
    for k in range(n):
        t = open_min + (k + 0.5) * span / n + offset_min
        t = max(open_min, min(close_min, t))
        out.append(f"{int(t) // 60:02d}:{int(t) % 60:02d}")
    return out


def main():
    c28 = json.load(open(os.path.join(SRC, "C28-21_Airport.geojson"), encoding="utf-8"))
    rp = json.load(open(os.path.join(SRC, "C28-21_AirportReferencePoint.geojson"), encoding="utf-8"))
    refpoints = {f["properties"]["C28_000"]: f["geometry"]["coordinates"] for f in rp["features"]}
    airports, services = [], []
    for ap in AIRPORTS:
        rw = runway_geometry(c28, refpoints, ap)
        st = stats(ap)
        airports.append(dict(id=ap["id"], name=ap["name"], coord=[round(rw["ref"][0], 6), round(rw["ref"][1], 6)],
                             heading=rw["heading"], runway=rw["length"], hours=rw["hours"],
                             polygon=rw["polygon"], stats=st,
                             routes=[dict(dest=d, perDay=n, intl=d in INTL) for d, n in ROUTES[ap["id"]]]))
        thr = rw["threshold"]
        rwy_end = move(thr, rw["heading"], rw["length"])
        name = ap["name"]
        for r, (dest, n) in enumerate(ROUTES[ap["id"]]):
            dc = (DEST[dest][1], DEST[dest][0])
            # 出発: 滑走路の端から離陸し、行き先の方角へ
            dep = climb_path(thr, rw["heading"], dc, rw["length"] / 1000 + 6)
            # 到着: 滑走路の反対側へ伸ばした最終進入コースから逆算し、向きを反転する
            back = climb_path(thr, (rw["heading"] + 180) % 360, dc, FINAL_KM)
            arr = back[::-1] + [move(thr, rw["heading"], 1500)]
            for kind, pts in (("dep", dep), ("arr", arr)):
                cum = [0.0]
                for a, b in zip(pts, pts[1:]):
                    cum.append(cum[-1] + meters(a, b))
                total = cum[-1]
                alt = []
                for d in cum:
                    if kind == "dep":
                        roll = rw["length"] * 0.7
                        alt.append(round(max(0.0, min(6000.0, (d - roll) * 0.08))))
                    else:
                        to_thr = total - 1500 - d
                        alt.append(round(max(0.0, min(6000.0, to_thr * 0.0524))))
                path = [[name if kind == "dep" else dest, [round(pts[0][0], 6), round(pts[0][1], 6)], 1]]
                path += [["", [round(x, 6), round(y, 6)], 0] for x, y in pts[1:-1]]
                path.append([dest if kind == "dep" else name, [round(pts[-1][0], 6), round(pts[-1][1], 6)], 1])
                # 到着便は、着陸の約 20 分前に 160km 圏へ入ってくる
                offset = r * 7 + (0 if kind == "dep" else 35)
                services.append(dict(
                    PLANE, id=f"air_{ap['id']}_{r}_{kind}", name=f"{name} ⇄ {dest}", group=GROUP["id"],
                    line=f"air_{ap['id']}", color=COLOR, loop=False, both=False, offset=0, path=path,
                    departures=schedule(n, rw["hours"], offset), flight=kind, airport=ap["id"],
                    fromName=name if kind == "dep" else dest, toName=dest if kind == "dep" else name,
                    intl=dest in INTL, alt=[[round(c), a] for c, a in zip(cum, alt)],
                    note="便数・時刻は推計"))
        print(f"{name}: 滑走路 {rw['heading']}° {rw['length']}m, {st['year']}年 乗降客 {st['passengers']:,} 人, "
              f"着陸 {st['landings']:,} 回, 推計 {sum(n for _, n in ROUTES[ap['id']])} 往復/日", file=sys.stderr)
    os.makedirs(OUT_DIR, exist_ok=True)
    out = dict(groups=[GROUP], replaces=[], lines=[], services=services, airports=airports,
               credit=None)
    with open(os.path.join(OUT_DIR, "airports.json"), "w", encoding="utf-8") as fp:
        json.dump(out, fp, ensure_ascii=False, separators=(",", ":"))


if __name__ == "__main__":
    main()
